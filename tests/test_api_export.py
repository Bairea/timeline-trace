import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture
def client(tmp_path, monkeypatch):
    import app.db as db
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setenv("APP_PASSWORD", "pw")
    import app.auth as auth
    auth._sessions.clear()
    from app.main import app
    with TestClient(app) as c:
        c.post("/api/login", json={"password": "pw"})
        yield c


def _seed_day(client, date="2026-09-20"):
    lines = ["06:35-07:02 早操", "07:22-07:32 早饭",
             "08:30-12:20 工作", "12:20-13:05 午饭与午休",
             "20:30-21:10 干活", "21:10-21:40 短视频",
             "22:40-05:55 睡觉"]
    for line in lines:
        r = client.post("/api/actual/quick", json={"date": date, "line": line})
        assert r.status_code == 200, r.text


def _parse_csv(content: bytes) -> list[dict]:
    text = content.decode("utf-8").lstrip("\ufeff")
    return list(csv.DictReader(io.StringIO(text)))


def test_export_csv_headers(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=csv")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]


def test_export_csv_has_bom(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=csv")
    assert r.content.decode("utf-8").startswith("\ufeff")


def test_export_actual_duration_sum_never_exceeds_a_day(client):
    """方案 B 的核心保证：实际明细的 duration_min 求和 <= 1440。"""
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20"
                   "&format=csv&scope=actual")
    rows = _parse_csv(r.content)
    total = sum(int(x["duration_min"]) for x in rows if x["duration_min"])
    assert total <= 1440, f"求和 {total} 分钟超过一天"
    assert total > 0


def test_export_actual_sum_equals_recorded_minutes(client):
    """明细求和精确等于记录的总时长（无重复、无遗漏）。"""
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20"
                   "&format=csv&scope=actual")
    rows = _parse_csv(r.content)
    total = sum(int(x["duration_min"]) for x in rows if x["duration_min"])
    expected = (27 + 10 + 230 + 45 + 40 + 30 + 435)
    assert total == expected


def test_export_actual_one_row_per_block(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20"
                   "&format=csv&scope=actual")
    rows = _parse_csv(r.content)
    assert len(rows) == 7      # 7 条记录，7 行


def test_export_both_has_template_section(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=csv")
    text = r.content.decode("utf-8").lstrip("\ufeff")
    assert "matched_template" in text
    assert "template_block" in text
    assert "unrecorded" in text


def test_export_both_template_section_has_18_blocks(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=csv")
    text = r.content.decode("utf-8").lstrip("\ufeff")
    lines = text.splitlines()
    # 找第二段表头之后的行
    idx = [i for i, l in enumerate(lines) if l.startswith("date,template_start")]
    assert idx, "应存在模板对照段表头"
    body = [l for l in lines[idx[0] + 1:] if l.strip()]
    assert len(body) == 18


def test_export_csv_escapes_commas(client):
    client.post("/api/actual/quick",
                json={"date": "2026-09-20", "line": "07:22-07:32 早饭,第二份"})
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=csv")
    assert '"早饭,第二份"' in r.content.decode("utf-8")


def test_export_xlsx_sheets(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["实际明细", "模板对照", "日汇总", "区间统计"]


def test_export_xlsx_actual_sum(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["实际明细"]
    total = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3]:
            total += int(row[3])
    assert 0 < total <= 1440


def test_export_cross_midnight_sleep_rendered(client):
    _seed_day(client)
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20"
                   "&format=csv&scope=actual")
    rows = _parse_csv(r.content)
    sleep = [x for x in rows if x["block_name"] == "睡觉"][0]
    assert sleep["start"] == "22:40"
    assert sleep["end"] == "05:55"
    assert sleep["duration_min"] == "435"


def test_export_rejects_bad_format(client):
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&format=pdf")
    assert r.status_code == 422


def test_export_rejects_bad_scope(client):
    r = client.get("/api/export?start=2026-09-20&end=2026-09-20&scope=nope")
    assert r.status_code == 422


def test_export_rejects_reversed_range(client):
    r = client.get("/api/export?start=2026-09-21&end=2026-09-20&format=csv")
    assert r.status_code == 422


def test_export_multi_day_range(client):
    _seed_day(client, "2026-09-20")
    _seed_day(client, "2026-09-21")
    r = client.get("/api/export?start=2026-09-20&end=2026-09-21"
                   "&format=csv&scope=actual")
    rows = _parse_csv(r.content)
    assert len(rows) == 14
    total = sum(int(x["duration_min"]) for x in rows if x["duration_min"])
    assert total <= 2 * 1440


def test_export_empty_range_still_valid(client):
    r = client.get("/api/export?start=2026-01-01&end=2026-01-02&format=csv")
    assert r.status_code == 200
    assert r.content.decode("utf-8").startswith("\ufeff")


def test_stats_endpoint(client):
    """项目推进时长按 category 统计，不再靠块名子串匹配。

    夹具里同时有「08:30-12:20 工作」（230 分钟）与「20:30-21:10 干活」
    （40 分钟），二者都落在模板的 category=项目推进 块内 → 270。
    旧口径只认名字含「干活」的块，会漏掉 230 分钟的「工作」。
    """
    _seed_day(client)
    r = client.get("/api/stats?start=2026-09-20&end=2026-09-20")
    assert r.status_code == 200
    body = r.json()
    assert len(body["daily"]) == 1
    metrics = dict(body["range"])
    assert metrics["days"] == 1
    assert metrics["project_min_total"] == 270


def test_stats_excludes_unrelated_block_inside_template_span(client):
    """模板块时段内做的别的事不应计入该模板块的 category。

    夹具里「21:10-21:40 短视频」落在模板「干活 20:30-22:00」区间内，
    但它不是项目推进。这是最容易让指标虚高的场景（干活时段摸鱼），
    故单列一条断言守住。
    """
    _seed_day(client)
    metrics = dict(client.get(
        "/api/stats?start=2026-09-20&end=2026-09-20").json()["range"])
    # 若短视频被误算，总量会是 300
    assert metrics["project_min_total"] == 270


def test_stats_requires_auth(tmp_path, monkeypatch):
    import app.db as db
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t2.db")
    monkeypatch.setenv("APP_PASSWORD", "pw")
    import app.auth as auth
    auth._sessions.clear()
    from app.main import app
    with TestClient(app) as c:
        assert c.get("/api/stats?start=2026-09-20&end=2026-09-20").status_code == 401
        assert c.get("/api/export?start=2026-09-20&end=2026-09-20").status_code == 401
