import io

import pytest
from openpyxl import load_workbook

from app.compare import BlockLike, compare_day
from app.export import (ACTUAL_HEADER, TEMPLATE_HEADER, export_csv,
                        export_csv_actual, export_xlsx, fmt_min,
                        rows_actual, rows_template)


# ---------- fmt_min ----------

def test_fmt_min_basic():
    assert fmt_min(390) == "06:30"
    assert fmt_min(0) == "00:00"
    assert fmt_min(1439) == "23:59"


def test_fmt_min_none_is_empty():
    assert fmt_min(None) == ""


def test_fmt_min_wraps_past_midnight():
    """跨零点必须还原为合法 HH:MM，不能在 Excel 里出现 30:30。"""
    assert fmt_min(1830) == "06:30"      # 1360(22:40) + 470
    assert fmt_min(1440) == "00:00"
    assert fmt_min(1500) == "01:00"


# ---------- rows_actual：实际明细，每块一行，可直接求和 ----------

def test_actual_rows_one_per_actual_block():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    assert len(rows) == 1
    assert rows[0][4] == "早操"


def test_actual_rows_no_duplicate_when_block_hits_two_templates():
    """一个实际块命中两个模板块时，明细表仍只出一行。

    这是方案 B 的核心：明细表可安全求和，不会超过 24 小时。
    """
    tpls = [BlockLike(440, 450, "早饭"), BlockLike(450, 455, "出门前准备")]
    acts = [BlockLike(442, 452, "早饭")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    assert len(rows) == 1
    assert rows[0][4] == "早饭"


def test_actual_rows_sum_never_exceeds_1440():
    """硬不变式：一天的实际块时长总和不能超过 1440 分钟。"""
    tpls = [BlockLike(440, 450, "早饭"), BlockLike(450, 455, "出门前准备"),
            BlockLike(390, 420, "早操")]
    acts = [BlockLike(442, 452, "早饭"), BlockLike(395, 422, "早操")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    total = sum(int(r[3]) for r in rows if r[3])
    assert total <= 1440, f"明细表求和 {total} 分钟超过一天"


def test_actual_rows_full_day_sum_is_exactly_recorded():
    """明细表求和必须精确等于实际记录的总时长。"""
    acts = [BlockLike(390, 420, "早操"), BlockLike(1230, 1270, "干活")]
    tpls = [BlockLike(390, 420, "早操")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    total = sum(int(r[3]) for r in rows if r[3])
    assert total == 30 + 40


def test_actual_rows_marks_unplanned():
    tpls = [BlockLike(1230, 1320, "干活")]
    acts = [BlockLike(100, 130, "夜里刷剧")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    r = [x for x in rows if x[4] == "夜里刷剧"][0]
    assert r[6] == "unplanned"
    assert r[5] == ""
    assert r[9] == "true"


def test_actual_rows_uses_actual_times():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    assert rows[0][1] == "06:35" and rows[0][2] == "07:02"
    assert rows[0][3] == "27"


def test_actual_rows_cross_midnight():
    tpls = [BlockLike(1360, 1830, "睡觉")]
    acts = [BlockLike(1370, 1840, "睡觉")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    assert rows[0][1] == "22:50"
    assert rows[0][2] == "06:40"
    assert rows[0][3] == "470"


def test_actual_rows_open_range_has_empty_duration():
    tpls = [BlockLike(1230, 1320, "干活")]
    acts = [BlockLike(1230, 1270, "干活")]
    rows = rows_actual("2026-09-20", compare_day(tpls, acts), acts)
    assert len(rows) == 1


# ---------- rows_template：模板对照，含 unrecorded ----------

def test_template_rows_one_per_template_block():
    tpls = [BlockLike(390, 420, "早操"), BlockLike(1320, 1340, "晚练")]
    acts = [BlockLike(395, 422, "早操")]
    rows = rows_template("2026-09-20", compare_day(tpls, acts))
    assert len(rows) == 2
    assert rows[0][4] == "早操"


def test_template_rows_include_unrecorded():
    tpls = [BlockLike(1340, 1360, "晚练")]
    rows = rows_template("2026-09-20", compare_day(tpls, []))
    assert rows[0][6] == "unrecorded"
    assert rows[0][1] == "22:20" and rows[0][2] == "22:40"


def test_template_rows_exclude_unplanned():
    """计划外是实际块的概念，不在模板对照表里出现。"""
    tpls = [BlockLike(1230, 1320, "干活")]
    acts = [BlockLike(100, 130, "夜里刷剧")]
    rows = rows_template("2026-09-20", compare_day(tpls, acts))
    assert all(r[6] != "unplanned" for r in rows)


def test_template_rows_deltas_for_aligned():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    rows = rows_template("2026-09-20", compare_day(tpls, acts))
    assert rows[0][7] == "5" and rows[0][8] == "-3"


# ---------- CSV ----------

def test_csv_has_utf8_bom():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    res = compare_day(tpls, acts)
    s = export_csv(rows_actual("2026-09-20", res, acts),
                   rows_template("2026-09-20", res))
    assert s[0] == "\ufeff"


def test_csv_actual_section_header():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    res = compare_day(tpls, acts)
    s = export_csv(rows_actual("2026-09-20", res, acts),
                   rows_template("2026-09-20", res)).lstrip("\ufeff")
    assert ",".join(ACTUAL_HEADER) in s
    assert ",".join(TEMPLATE_HEADER) in s


def test_csv_escapes_commas_in_names():
    """块名含逗号时必须加引号，否则列会错位。"""
    tpls = [BlockLike(390, 420, "早操,第二组")]
    acts = [BlockLike(395, 422, "早操,第二组")]
    res = compare_day(tpls, acts)
    s = export_csv(rows_actual("2026-09-20", res, acts),
                   rows_template("2026-09-20", res)).lstrip("\ufeff")
    assert '"早操,第二组"' in s


def test_csv_sections_are_separated_by_blank_line():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    res = compare_day(tpls, acts)
    s = export_csv(rows_actual("2026-09-20", res, acts),
                   rows_template("2026-09-20", res)).lstrip("\ufeff")
    lines = s.splitlines()
    blank = [i for i, l in enumerate(lines) if l.strip() == ""]
    assert blank, "两个 section 之间应有空行分隔"


# ---------- XLSX ----------

def test_xlsx_has_sheets():
    tpls = [BlockLike(390, 420, "早操")]
    acts = [BlockLike(395, 422, "早操")]
    res = compare_day(tpls, acts)
    data = export_xlsx(rows_actual("2026-09-20", res, acts),
                       rows_template("2026-09-20", res),
                       [["2026-09-20", 27, 0, 1, 0, 0, 0]],
                       [["days", 1]])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["实际明细", "模板对照", "日汇总", "区间统计"]


def test_xlsx_actual_sheet_row_count():
    tpls = [BlockLike(440, 450, "早饭"), BlockLike(450, 455, "出门前准备")]
    acts = [BlockLike(442, 452, "早饭")]
    res = compare_day(tpls, acts)
    data = export_xlsx(rows_actual("2026-09-20", res, acts),
                       rows_template("2026-09-20", res))
    wb = load_workbook(io.BytesIO(data))
    assert wb["实际明细"].max_row == 2          # 表头 + 1 行
    assert wb["模板对照"].max_row == 3          # 表头 + 2 行


def test_xlsx_empty_is_valid():
    data = export_xlsx([], [])
    wb = load_workbook(io.BytesIO(data))
    assert wb["实际明细"].max_row == 1
    assert wb["模板对照"].max_row == 1
